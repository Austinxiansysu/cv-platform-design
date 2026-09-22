from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/xianjianhua/Desktop/CV platform design/JD分析练习稿.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=11, bold=False, color="000000"):
    run.font.name = "Source Han Serif SC VF"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Source Han Serif SC VF")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Source Han Serif SC VF")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Source Han Serif SC VF")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Source Han Serif SC VF")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    styles = doc.styles
    for name, size, bold, before, after in [
        ("Normal", 11, False, 0, 6),
        ("Title", 24, True, 0, 16),
        ("Subtitle", 11, False, 0, 14),
        ("Heading 1", 17, True, 16, 8),
        ("Heading 2", 13, True, 12, 6),
        ("Heading 3", 11, True, 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Source Han Serif SC VF"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Source Han Serif SC VF")
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Source Han Serif SC VF")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Source Han Serif SC VF")
        style._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Source Han Serif SC VF")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)


def add_meta_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.55)
    for i, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Inches(1.35)
        cells[1].width = Inches(5.55)
        for cell in cells:
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], "DCE6F1")
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        set_run_font(p0.add_run(label), size=10.5, bold=True)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(value), size=10.5)
    doc.add_paragraph()


def add_numbered_items(doc, items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{idx}  {item}")
        set_run_font(r)


def add_analysis_form(doc):
    doc.add_heading("我的分析", level=2)
    intro = doc.add_paragraph("请先独立完成。重点区分 JD 明确写出的事实、自己的推断和仍然未知的信息。")
    intro.paragraph_format.space_after = Pt(9)

    questions = [
        "1 只看岗位名称时，我原来认为它在做什么",
        "2 读完 JD 后，我认为最核心的三项日常任务是什么",
        "3 上述判断中，哪些是 JD 明确陈述，哪些是我的推断",
        "4 我最想做其中哪项任务，为什么",
        "5 我最不想做其中哪项任务，为什么",
        "6 我有哪些真实经历或能力证据可以支持这份工作",
        "7 我目前缺少什么，包括能力缺口和证据缺口",
        "8 哪些信息无法从 JD 判断，需要在面试或沟通中确认",
        "9 如果完成三个月实习，它最可能帮助我验证什么职业假设",
    ]
    for q in questions:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(q), size=10.5, bold=True)
        ans = doc.add_paragraph("在此填写")
        ans.paragraph_format.left_indent = Inches(0.25)
        ans.paragraph_format.space_after = Pt(10)
        set_run_font(ans.runs[0], size=10.5, color="7F7F7F")

    doc.add_heading("初步结论", level=3)
    for label, options in [
        ("想做程度", "□ 高    □ 中    □ 低"),
        ("当前准备度", "□ 高    □ 中    □ 低"),
        ("现实可行性", "□ 可行    □ 待确认    □ 明确冲突"),
        ("是否值得申请或继续了解", "□ 值得申请    □ 先补充信息    □ 暂不申请"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(f"{label}  "), size=10.5, bold=True)
        set_run_font(p.add_run(options), size=10.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    set_run_font(p.add_run("最终理由"), size=10.5, bold=True)
    ans = doc.add_paragraph("在此填写")
    ans.paragraph_format.left_indent = Inches(0.25)
    set_run_font(ans.runs[0], size=10.5, color="7F7F7F")


def add_job(doc, title, meta, duties, requirements, notes=None):
    doc.add_heading(title, level=1)
    add_meta_table(doc, meta)
    doc.add_heading("岗位职责原文", level=2)
    add_numbered_items(doc, duties)
    doc.add_heading("岗位要求原文", level=2)
    add_numbered_items(doc, requirements)
    if notes:
        doc.add_heading("已知边界", level=2)
        for note in notes:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            set_run_font(p.add_run(note))
    add_analysis_form(doc)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
configure_styles(doc)

title = doc.add_paragraph(style="Title")
set_run_font(title.add_run("三条岗位 JD 分析练习稿"), size=24, bold=True)
subtitle = doc.add_paragraph(style="Subtitle")
set_run_font(subtitle.add_run("用于职业任务理解与个人匹配判断"), size=11)

p = doc.add_paragraph(
    "这份练习稿包含三条具有代表性的真实岗位：新型 AI 交叉岗位、典型商业分析岗位，以及名称与实际任务容易产生偏差的投后管理岗位。请先根据 JD 独立分析，不追求得出唯一正确答案，而是练习区分想做程度、当前准备度和现实可行性。"
)
p.paragraph_format.space_after = Pt(10)

doc.add_heading("填写原则", level=2)
for text in [
    "事实只来自 JD 原文，不根据公司名或岗位名补充想象。",
    "JD 没写的信息标记为未知，不用猜测代替。",
    "能力判断必须对应自己的真实课程、项目或经历证据。",
    "想做程度和当前准备度分别判断，不用一个分数混在一起。",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(text))

doc.add_page_break()

add_job(
    doc,
    "岗位一 字节跳动飞书 FDE 实习生",
    [
        ("公司与团队", "字节跳动 飞书商业化"),
        ("地点", "广州"),
        ("招聘项目", "ByteIntern"),
        ("年级学历", "2027 届本科及以上学历在读"),
        ("官网职类", "销售  销售支持"),
        ("来源", "字节跳动校园招聘官网  职位 ID A258844B"),
    ],
    [
        "走进企业客户的真实工作场景，理解业务难题，使用飞书 AI 工具搭建能够解决问题的智能工作流方案。",
        "配合商务或渠道经理开展制造、零售、电商、餐饮等行业的客户调研，识别低效环节、业务痛点和 AI 落地机会。",
        "运用 AI 智能助手、自动化工作流、多维表格等无代码或低代码能力，配置和搭建可演示或可投入使用的解决方案 Demo。",
        "将成功案例总结为标准化使用指南 SOP 和演示模板，供销售团队或合作伙伴复用。",
        "以培训或技术支持形式向代理商和生态合作伙伴传递 AI 方案搭建经验。",
    ],
    [
        "2027 届本科及以上学历在读，专业不限；计算机或信息技术与商科或管理类复合背景优先。",
        "对 AI 和大模型等新技术有兴趣，愿意钻研和实际操作软件工具；具备基础逻辑判断、数据结构或无代码开发经验者优先。",
        "能够拆解复杂业务问题，并向客户或非技术人员清楚解释技术方案。",
        "面对未知场景保持好奇心，主动寻找解决方案，并能在持续试错中推进工作。",
    ],
    [
        "岗位使用 AI 并搭建方案，但官网将其归为销售支持，不能直接等同于 AI 产品经理。",
        "JD 没有明确每周到岗天数和具体实习时长。",
    ],
)

doc.add_page_break()

add_job(
    doc,
    "岗位二 腾讯商业分析实习",
    [
        ("公司", "腾讯"),
        ("招聘类型", "应届实习"),
        ("面试城市", "远程面试"),
        ("涉及事业群", "CSIG  IEG  PCG  WXG"),
        ("工作地点", "岗位页面未明确具体城市"),
        ("来源", "腾讯校园招聘官网"),
    ],
    [
        "关注行业发展趋势，运用商业分析方法深入分析行业商业模式、产品创新、市场和竞争环境变化。",
        "通过定性和定量研究洞察用户行为变化及需求演进，为公司业务寻找战略机会。",
        "为各层级业务管理和商业决策提供专项战略分析及决策建议。",
        "协助制定事业群和部门的商业计划及项目计划。",
        "支持战略级跨团队、跨事业群项目，协助推进执行与落地。",
    ],
    [
        "学业成绩优异，快速学习能力强，保持好奇心并对互联网行业有热情。",
        "具备领导力、责任感、正直、奉献和进取精神，以及团队合作能力。",
        "沟通和人际交往能力突出，熟练使用英文。",
        "理解基本商业问题，具备结构化思维，能够拆解和解决问题，并具有一定商业意识。",
    ],
    [
        "JD 提到定性与定量研究，但没有明确 SQL、Python 或其他具体数据工具要求。",
        "页面没有说明不同事业群的具体岗位内容是否完全一致。",
    ],
)

doc.add_page_break()

add_job(
    doc,
    "岗位三 投后管理实习生",
    [
        ("公司", "群内截图未显示"),
        ("地点", "广州"),
        ("学历", "本科及以上在读  硕士在读优先"),
        ("实习时长", "6 个月及以上"),
        ("出勤要求", "每周 4 天及以上"),
        ("来源", "中大岭院实习群转发  非官网来源"),
    ],
    [
        "协助投后管理经理开展投后运营、融资支撑、基金募资辅助及合规风控相关工作。",
        "跟踪被投或孵化项目数据，整理经营分析材料和档案台账，辅助股东沟通及风险预警基础工作。",
        "推进股权和基金融资相关事项，包括 BP、PPM 等融资材料整理、尽调资料汇总和数据核对。",
        "维护投资人库和 LP 基础信息，配合基金募资文案及路演材料整理。",
        "开展财务、法务和税务合规基础工作，包括合同整理、资料归档、政策信息收集和融资尽调答疑支持。",
        "完成 PR 和 GR 基础工作，包括舆情监控、政策申报资料整理和品牌宣传辅助事项。",
        "配合经理完成其他投后管理辅助工作，确保工作有序推进。",
    ],
    [
        "本科及以上在读，硕士在读优先；财务、投资、金融、经济等相关专业，海外留学经历或海外院校在读优先。",
        "能够保证 6 个月及以上全职实习，每周出勤 4 天及以上。",
        "具有投行、会计师事务所、律师事务所或股权投资机构实习经验者优先。",
        "具备基础财务分析能力，熟悉 Excel、Word 和 PPT。",
        "工作严谨细致、执行力强，具备沟通协调和学习能力，并对投后管理和股权投资有兴趣。",
    ],
    [
        "公司名称没有出现在截图中，实际投递前需要核验发布主体和岗位真实性。",
        "该岗位同时包含投后运营、募资、合规、材料和宣传工作，不能只根据投资标签理解岗位。",
    ],
)

doc.core_properties.title = "三条岗位 JD 分析练习稿"
doc.core_properties.subject = "职业任务理解与个人匹配判断"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.save(OUT)
print(OUT)

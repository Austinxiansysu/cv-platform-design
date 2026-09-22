from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SRC = Path("/Users/xianjianhua/Desktop/CV platform design/JD分析练习稿.docx")
OUT = Path("/Users/xianjianhua/Desktop/CV platform design/JD分析练习稿_导师修订.docx")


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


doc = Document(SRC)


def section_bounds(title):
    starts = [i for i, p in enumerate(doc.paragraphs) if p.text.strip() == title]
    if not starts:
        raise ValueError(f"Section not found: {title}")
    start = starts[0]
    end = len(doc.paragraphs)
    for i in range(start + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].style.name == "Heading 1":
            end = i
            break
    return start, end


def replace_answer(title, question_prefix, answer):
    start, end = section_bounds(title)
    paras = doc.paragraphs
    for i in range(start, end):
        if paras[i].text.strip().startswith(question_prefix):
            j = i + 1
            while j < end and not paras[j].text.strip():
                j += 1
            if j >= end:
                raise ValueError(f"Answer not found: {title} / {question_prefix}")
            paras[j].text = answer
            return
    raise ValueError(f"Question not found: {title} / {question_prefix}")


def replace_line(title, prefix, new_text):
    start, end = section_bounds(title)
    for i in range(start, end):
        if doc.paragraphs[i].text.strip().startswith(prefix):
            doc.paragraphs[i].text = new_text
            return doc.paragraphs[i]
    raise ValueError(f"Line not found: {title} / {prefix}")


def set_final_reason(title, reason):
    start, end = section_bounds(title)
    paras = doc.paragraphs
    for i in range(start, end):
        if paras[i].text.strip() == "最终理由":
            j = i + 1
            while j < end and not paras[j].text.strip():
                j += 1
            if j < end:
                paras[j].text = reason
            else:
                insert_paragraph_after(paras[i], reason, "Normal")
            return
    decision = replace_line(title, "是否值得申请或继续了解", next(
        p.text for p in paras[start:end] if p.text.strip().startswith("是否值得申请或继续了解")
    ))
    heading = insert_paragraph_after(decision, "最终理由", "Heading 3")
    insert_paragraph_after(heading, reason, "Normal")


fde = "岗位一 字节跳动飞书 FDE 实习生"
replace_answer(fde, "2 读完 JD", "调研企业客户的业务流程和痛点并识别 AI 落地机会；使用飞书 AI 与低代码工具搭建解决方案 Demo；把成功案例沉淀为 SOP、模板和培训材料供销售及合作伙伴复用。")
replace_answer(fde, "3 上述判断中", "JD 明确陈述了客户调研、AI 方案搭建、SOP 沉淀和伙伴培训。我的推断是实习生可能拥有较完整的方案搭建机会，以及这些任务能直接对应 AI 产品经理工作。")
replace_answer(fde, "4 我最想做", "最想做业务痛点分析和 AI 方案搭建。我已有 AI 辅助完成企业网站的经历，喜欢从真实需求出发快速形成可用成果，也希望继续验证自己是否适合业务与技术交叉工作。")
replace_answer(fde, "5 我最不想做", "目前最不确定的是代理商培训、持续技术支持和可能重复配置相似方案。这些任务可能包含较多外部沟通和支持性工作，我尚未验证自己是否长期喜欢。")
replace_answer(fde, "6 我有哪些", "企业官网项目能够证明我使用 AI 工具完成从需求理解到交付的基本能力，也能证明我愿意亲自动手搭建工作流；但它还不能证明企业客户调研、飞书低代码配置、SOP 编写或培训能力。")
replace_answer(fde, "7 我目前缺少", "缺少企业客户调研、真实产品运营场景、飞书低代码工具、SOP 与培训交付的直接经验；现有网站项目的过程和结果也需要整理成可核验的证据。")
replace_answer(fde, "8 哪些信息", "需要确认每周到岗天数、实习时长、客户调研与方案搭建的时间占比、实习生可以独立负责的范围、培训支持工作的占比和实际工作强度。")
replace_answer(fde, "9 如果完成", "验证我是否喜欢面向企业客户发现业务问题，并利用 AI 和低代码工具实施解决方案；进一步区分我更适合 AI 产品、解决方案顾问、FDE、售前支持还是客户成功。")
replace_line(fde, "当前准备度", "当前准备度  □ 高    ☑ 中    □ 低")
replace_line(fde, "现实可行性", "现实可行性  □ 可行    □ 待确认    ☑ 明确冲突")
replace_line(fde, "是否值得申请或继续了解", "是否值得申请或继续了解  □ 值得申请    □ 先补充信息    ☑ 暂不申请")
set_final_reason(fde, "任务兴趣较高，也有 AI 辅助开发的间接证据，但岗位实质包含客户调研、方案实施和销售支持，相关能力仍待验证。按正常四年制学制预计毕业年份与岗位要求的 2027 届不符，因此当前不申请，把它保留为未来职业方向样本。")


tencent = "岗位二 腾讯商业分析实习"
replace_answer(tencent, "2 读完 JD", "研究行业趋势、商业模式、产品创新和竞争环境；通过定性与定量研究洞察用户行为和需求；形成战略分析与决策建议，并协助商业计划和跨团队项目落地。")
replace_answer(tencent, "3 上述判断中", "JD 明确陈述了行业与竞争研究、用户洞察、战略建议、商业计划和跨团队落地。我的推断是实习生会大量进行定量分析，且产出能够被其他团队直接复用并产生可观察效果。")
replace_answer(tencent, "4 我最想做", "最想做行业商业模式、外部环境和用户需求分析，并把分析转化为决策建议。这类任务同时包含商业理解、结构化分析和成果输出，与我的探索方向较一致。")
replace_answer(tencent, "5 我最不想做", "目前最不确定的是跨事业群项目推进。如果工作主要变成会议、流程协调和跟进执行，而分析深度较低，我的兴趣可能下降。")
replace_answer(tencent, "6 我有哪些", "岭南杯案例分析可以证明我有商业问题拆解和团队报告经验，数学建模比赛可以证明定量分析与结构化思考能力；但尚不能证明真实互联网业务、用户研究或战略落地经验。")
replace_answer(tencent, "7 我目前缺少", "缺少互联网商业模式、用户研究、战略分析流程和真实公司数据的经验，也缺少能够证明建议被业务采用并产生效果的成果证据。")
replace_answer(tencent, "8 哪些信息", "需要确认具体事业群和团队、工作地点与毕业年份要求、定性和定量研究的比例、常用数据工具、分析成果的主要使用者、实习生的项目责任范围以及实际工作强度。")
replace_answer(tencent, "9 如果完成", "验证我是否真正喜欢商业分析的日常研究与项目推进，并理解互联网公司的商业分析与传统金融研究在对象、节奏和成果使用方式上的差异。")
replace_line(tencent, "现实可行性", "现实可行性  □ 可行    ☑ 待确认    □ 明确冲突")
replace_line(tencent, "是否值得申请或继续了解", "是否值得申请或继续了解  □ 值得申请    ☑ 先补充信息    □ 暂不申请")
set_final_reason(tencent, "任务方向与我对商业、数据和产业研究的兴趣较契合，比赛经历能提供部分结构化分析证据；但具体事业群、毕业年份、工作地点、定量研究占比和实习生职责仍不清楚，应先补充关键信息再决定是否申请。")


post = "岗位三 投后管理实习生"
replace_answer(post, "2 读完 JD", "第一，跟踪被投项目经营情况并协助风险预警；第二，整理融资、募资和路演材料并维护投资人信息；第三，完成合同、合规、政策申报、档案和宣传等运营支持工作。")
replace_answer(post, "3 上述判断中", "JD 明确陈述了项目数据跟踪、融资募资材料、投资人库、合规归档和 PR、GR 支持。我的推断是材料整理可能占据较大比例，实习生的独立分析空间可能有限。")
replace_answer(post, "4 我最想做", "相对最想参与被投项目经营分析、融资数据核对和风险预警，因为这些任务包含一定的金融分析和判断，而不仅是流程处理。")
replace_answer(post, "5 我最不想做", "最不想长期承担档案台账、合同归档、路演材料和宣传辅助等重复性整理工作；对 PR、GR 也缺少兴趣和了解。")
replace_answer(post, "6 我有哪些", "目前没有直接的投后管理、融资尽调或合规经验。案例分析和数学建模只能间接证明基础分析能力，不能充分证明岗位胜任度。")
replace_answer(post, "7 我目前缺少", "缺少投融资、财务分析、风险管控、基金募资、合规和尽调知识，也缺少 Excel 财务分析及真实股权投资项目的证据。")
replace_answer(post, "8 哪些信息", "需要确认各类任务的时间占比、实习生能否独立负责经营分析、是否参与正式风险讨论、主要汇报对象、工作强度，以及材料整理是否构成日常工作的主体。")
replace_answer(post, "9 如果完成", "验证我是否对股权投资中的投后经营分析和风险控制真正感兴趣，以及自己能否接受投后工作中大量材料、合规与协调任务。")
set_final_reason(post, "岗位包含大量材料、募资、合规和协调工作，与我偏好的问题分析和成果落地不完全一致；当前直接能力证据不足，且 6 个月、每周 4 天的条件可能与课程冲突，因此暂不申请，仅保留为了解股权投资投后链条的样本。")


doc.core_properties.title = "三条岗位 JD 分析练习稿 导师修订版"
doc.core_properties.subject = "职业任务理解与个人匹配判断"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.save(OUT)
print(OUT)
